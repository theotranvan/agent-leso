"""Tests du générateur de géométrie structurelle paramétrique."""
from __future__ import annotations

import pytest

from app.services.structure.geometry_builder import (
    FrameParams,
    build_frame_geometry,
    geometry_summary,
)


class TestBuildFrameGeometry:
    def test_comptes_portique_3_niveaux(self) -> None:
        # « Les Tilleuls » : 3 niveaux, trame 3×2.
        p = FrameParams(n_levels=3, n_bays_x=3, n_bays_y=2)
        geo = build_frame_geometry(p)
        s = geometry_summary(geo)

        nx, ny = 4, 3  # (bays+1)
        # Nœuds : (niveaux+1) plans × nx × ny
        assert s["nb_nodes"] == (3 + 1) * nx * ny  # 48
        # Poteaux : niveaux × nx × ny
        assert s["nb_columns"] == 3 * nx * ny       # 36
        # Poutres : par plancher (3), X = ny*n_bays_x=3*3=9, Y = nx*n_bays_y=4*2=8 → 17×3
        assert s["nb_beams"] == 3 * (ny * 3 + nx * 2)  # 51
        # Appuis : un par poteau de base
        assert s["nb_supports"] == nx * ny          # 12

    def test_appuis_a_la_base_uniquement(self) -> None:
        geo = build_frame_geometry(FrameParams(n_levels=2, n_bays_x=1, n_bays_y=1))
        base_nodes = {n["id"] for n in geo["nodes"] if n["z"] == 0.0}
        support_nodes = {s["node"] for s in geo["supports"]}
        assert support_nodes == base_nodes
        assert all(s["type"] == "fixed" for s in geo["supports"])

    def test_coordonnees_trame(self) -> None:
        p = FrameParams(n_levels=1, n_bays_x=2, bay_x_m=5.0, n_bays_y=1, bay_y_m=6.0, story_height_m=3.0)
        geo = build_frame_geometry(p)
        xs = sorted({n["x"] for n in geo["nodes"]})
        ys = sorted({n["y"] for n in geo["nodes"]})
        zs = sorted({n["z"] for n in geo["nodes"]})
        assert xs == [0.0, 5.0, 10.0]
        assert ys == [0.0, 6.0]
        assert zs == [0.0, 3.0]

    def test_ids_nodes_uniques_et_references_valides(self) -> None:
        geo = build_frame_geometry(FrameParams(n_levels=3, n_bays_x=3, n_bays_y=2))
        node_ids = {n["id"] for n in geo["nodes"]}
        assert len(node_ids) == len(geo["nodes"])  # pas de doublon
        # Chaque membre référence des nœuds existants
        for m in geo["members"]:
            assert m["node_start"] in node_ids
            assert m["node_end"] in node_ids
            assert m["node_start"] != m["node_end"]

    def test_section_et_materiau_propages(self) -> None:
        p = FrameParams(column_section="POT_40x40", beam_section="POU_30x60", material="C25/30")
        geo = build_frame_geometry(p)
        cols = [m for m in geo["members"] if m["type"] == "column"]
        beams = [m for m in geo["members"] if m["type"] == "beam"]
        assert all(c["section"] == "POT_40x40" and c["material"] == "C25/30" for c in cols)
        assert all(b["section"] == "POU_30x60" and b["material"] == "C25/30" for b in beams)

    @pytest.mark.parametrize("bad", [
        FrameParams(n_levels=0),
        FrameParams(n_bays_x=0),
        FrameParams(bay_x_m=0),
        FrameParams(story_height_m=-1),
    ])
    def test_params_invalides_levent_valueerror(self, bad: FrameParams) -> None:
        with pytest.raises(ValueError):
            build_frame_geometry(bad)


class TestGeometryProducesValidSaf:
    def test_saf_genere_avec_geometrie(self) -> None:
        """La géométrie générée doit produire un SAF xlsx non vide avec nœuds + membres."""
        from app.services.structure.saf_generator import generate_saf_xlsx

        geo = build_frame_geometry(FrameParams(n_levels=2, n_bays_x=2, n_bays_y=2))
        model = {
            "project_info": {
                "referentiel": "sia",
                "exposure_class": "XC2",
                "consequence_class": "CC2",
                "seismic_zone": "Z1b",
            },
            **geo,
        }
        xlsx = generate_saf_xlsx(model)
        assert isinstance(xlsx, bytes)
        assert len(xlsx) > 2000  # un xlsx réel fait plusieurs Ko

        # Relecture : la feuille des membres contient bien nos éléments
        import io

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx))
        members_ws = wb["Structural1DMembers"]
        rows = list(members_ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == len(geo["members"])
        nodes_ws = wb["StructuralNodes"]
        node_rows = list(nodes_ws.iter_rows(min_row=2, values_only=True))
        assert len(node_rows) == len(geo["nodes"])
