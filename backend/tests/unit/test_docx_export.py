"""Export Word (.docx) — fidélité et robustesse du générateur.

L'export Word est un livrable que l'ingénieur OUVRE et RETRAVAILLE dans Word.
Un .docx corrompu = retour client immédiat. Ces tests vérifient que :
  - le fichier produit est un OOXML valide (relisible par python-docx)
  - les structures sémantiques (titres, listes, tableaux, gras) sont rendues
  - le HTML réel produit par les agents passe sans planter
  - un HTML malformé / vide ne fait jamais planter (fail-soft)
  - le branding (couleur primaire) est appliqué
"""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.services.docx_generator import html_to_docx_bytes, text_to_docx_bytes

# En-tête ZIP d'un OOXML (.docx = archive ZIP) : "PK\x03\x04".
ZIP_MAGIC = b"PK\x03\x04"


def _reopen(data: bytes) -> Document:
    """Relit les octets .docx → preuve que le fichier est un OOXML valide."""
    return Document(BytesIO(data))


class TestValidDocx:
    def test_returns_valid_ooxml_bytes(self):
        data = html_to_docx_bytes("<p>Bonjour</p>", title="Test")
        assert isinstance(data, bytes) and len(data) > 0
        assert data[:4] == ZIP_MAGIC, "Le .docx n'est pas une archive ZIP valide"
        doc = _reopen(data)  # ne doit pas lever
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Bonjour" in text
        assert "Test" in text  # le titre

    def test_text_fallback_valid(self):
        data = text_to_docx_bytes("Para 1\n\nPara 2", title="Repli")
        assert data[:4] == ZIP_MAGIC
        doc = _reopen(data)
        joined = "\n".join(p.text for p in doc.paragraphs)
        assert "Para 1" in joined and "Para 2" in joined


class TestStructuralFidelity:
    def test_headings_become_headings(self):
        html = "<h2>Section A</h2><p>texte</p><h3>Sous-section</h3>"
        doc = _reopen(html_to_docx_bytes(html))
        styles = [p.style.name for p in doc.paragraphs]
        assert any("Heading" in s for s in styles), f"Aucun titre détecté : {styles}"

    def test_bullet_and_numbered_lists(self):
        html = "<ul><li>puce un</li><li>puce deux</li></ul><ol><li>num un</li></ol>"
        doc = _reopen(html_to_docx_bytes(html))
        styles = [p.style.name for p in doc.paragraphs]
        assert any("List Bullet" in s for s in styles), f"Pas de liste à puces : {styles}"
        assert any("List Number" in s for s in styles), f"Pas de liste numérotée : {styles}"
        texts = [p.text for p in doc.paragraphs]
        assert "puce un" in texts and "num un" in texts

    def test_table_rendered_with_all_cells(self):
        html = (
            "<table>"
            "<tr><th>Désignation</th><th>Prix</th></tr>"
            "<tr><td>Radiateur</td><td>250</td></tr>"
            "<tr><td>Vanne</td><td>40</td></tr>"
            "</table>"
        )
        doc = _reopen(html_to_docx_bytes(html))
        assert len(doc.tables) == 1, "Tableau non rendu"
        table = doc.tables[0]
        # 3 lignes (en-tête + 2), 2 colonnes
        assert len(table.rows) == 3
        assert len(table.columns) == 2
        cells = [c.text for row in table.rows for c in row.cells]
        for expected in ("Désignation", "Prix", "Radiateur", "250", "Vanne", "40"):
            assert expected in cells, f"Cellule manquante : {expected}"

    def test_bold_run_preserved(self):
        doc = _reopen(html_to_docx_bytes("<p>Texte <strong>important</strong> ici</p>"))
        bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
        assert any("important" in t for t in bold_runs), "Le gras n'est pas conservé"

    def test_branding_color_applied_to_title(self):
        data = html_to_docx_bytes(
            "<p>corps</p>", title="Mon Bureau",
            branding={"primary_color": "#1A73E8"},
        )
        doc = _reopen(data)
        # le titre (heading 0) doit porter la couleur primaire sur au moins un run
        colored = []
        for p in doc.paragraphs:
            for r in p.runs:
                rgb = r.font.color and r.font.color.rgb
                if rgb is not None:
                    colored.append(str(rgb))
        assert any(c.upper() == "1A73E8" for c in colored), \
            f"Couleur de marque non appliquée : {colored}"

    def test_cover_page_fields(self):
        data = html_to_docx_bytes(
            "<p>corps</p>", title="CCTP Chauffage", subtitle="Lot 230",
            project_info={"Projet": "Villa Dupont", "Auteur": "Ing. X"},
        )
        doc = _reopen(data)
        joined = "\n".join(p.text for p in doc.paragraphs)
        # Les infos projet sont rendues dans un tableau de garde (à la charte).
        joined += "\n" + "\n".join(
            c.text for t in doc.tables for r in t.rows for c in r.cells
        )
        assert "Villa Dupont" in joined and "Ing. X" in joined
        assert "Lot 230" in joined


class TestRobustness:
    def test_empty_html_does_not_crash(self):
        data = html_to_docx_bytes("", title="Vide")
        assert data[:4] == ZIP_MAGIC
        _reopen(data)

    def test_malformed_html_does_not_crash(self):
        # balises non fermées, imbrication illégale, tableau cassé
        bad = "<p>début <strong>gras <ul><li>item</p></strong><table><tr><td>x"
        data = html_to_docx_bytes(bad, title="Malformé")
        assert data[:4] == ZIP_MAGIC
        doc = _reopen(data)
        assert "item" in "\n".join(p.text for p in doc.paragraphs) or doc.tables

    def test_unknown_tags_ignored_gracefully(self):
        html = "<section><article><p>contenu</p><video src='x'></video></article></section>"
        doc = _reopen(html_to_docx_bytes(html))
        assert "contenu" in "\n".join(p.text for p in doc.paragraphs)

    def test_html_entities_decoded(self):
        doc = _reopen(html_to_docx_bytes("<p>l&apos;eau &amp; le feu &lt;3</p>"))
        joined = "\n".join(p.text for p in doc.paragraphs)
        assert "l'eau & le feu <3" in joined

    @pytest.mark.parametrize("html", [
        "<h2>Titre</h2><p>Para avec <em>italique</em> et <strong>gras</strong>.</p>"
        "<ul><li>a</li><li>b</li></ul>"
        "<table><tr><th>K</th><th>V</th></tr><tr><td>1</td><td>2</td></tr></table>",
        "<h1>Rapport</h1><p>" + "Lorem ipsum dolor sit amet. " * 200 + "</p>",
    ])
    def test_representative_agent_html(self, html):
        """HTML représentatif de ce que produisent cctp/coordination/rapport."""
        data = html_to_docx_bytes(html, title="Livrable")
        assert data[:4] == ZIP_MAGIC
        _reopen(data)
